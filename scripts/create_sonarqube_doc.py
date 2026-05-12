from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ACCENT = RGBColor(28, 63, 170)
TEXT = RGBColor(40, 40, 40)
MUTED = RGBColor(95, 95, 95)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def style_run(run, *, size=11, bold=False, color=TEXT, font="Aptos"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_paragraph(doc, text="", *, style=None, align=None, space_after=6, space_before=0):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        style_run(r)
    if align is not None:
        p.alignment = align
    fmt = p.paragraph_format
    fmt.space_after = Pt(space_after)
    fmt.space_before = Pt(space_before)
    fmt.line_spacing = 1.15
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    style_run(r)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.1
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    style_run(r)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.1
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    r = p.add_run(text)
    style_run(r, size=16 if level == 1 else 13, bold=True, color=ACCENT)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_info_table(doc):
    table = doc.add_table(rows=4, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True

    data = [
        ("Proyecto", "Mi NextJS App"),
        ("Ruta analizada", "/Users/greivin/Documents/GitHub/Edu360/edu360_web"),
        ("Servidor SonarQube", "http://127.0.0.1:9000"),
        ("Tecnología de despliegue", "Docker con imagen sonarqube:community"),
    ]

    for row_idx, (label, value) in enumerate(data):
        for col_idx, text in enumerate((label, value)):
            cell = table.cell(row_idx, col_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            style_run(r, bold=(col_idx == 0))
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            if col_idx == 0:
                set_cell_shading(cell, "EAF0FF")
    return table


def add_command_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6FA")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    for idx, line in enumerate(lines):
        r = p.add_run(line)
        style_run(r, size=10, color=RGBColor(20, 20, 20), font="Courier New")
        if idx < len(lines) - 1:
            r.add_break()
    doc.add_paragraph()


def build_doc(path):
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Instalación y Acceso a SonarQube")
    style_run(r, size=22, bold=True, color=ACCENT)
    title.paragraph_format.space_after = Pt(8)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Guía paso a paso para el proyecto Edu360")
    style_run(r, size=12, color=MUTED)
    subtitle.paragraph_format.space_after = Pt(18)

    add_info_table(doc)
    doc.add_paragraph()

    add_paragraph(
        doc,
        "Este documento describe la instalación local de SonarQube, la ejecución del entorno y los pasos necesarios para acceder al dashboard del proyecto analizado.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=10,
    )

    add_heading(doc, "1. Requisitos previos", level=1)
    add_bullet(doc, "Tener Docker Desktop instalado y en ejecución.")
    add_bullet(doc, "Contar con el proyecto local clonado en la máquina.")
    add_bullet(doc, "Tener Node.js y las dependencias del proyecto instaladas en la carpeta web.")
    add_bullet(doc, "Disponer de un navegador web para ingresar al dashboard de SonarQube.")

    add_heading(doc, "2. Instalación y levantamiento de SonarQube", level=1)
    add_paragraph(
        doc,
        "En este proyecto, SonarQube se ejecuta mediante Docker usando un archivo docker-compose. El procedimiento fue el siguiente:",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_number(doc, "Ubicar la carpeta del frontend del proyecto: edu360_web.")
    add_number(doc, "Verificar la existencia del archivo docker-compose.yml con el servicio de SonarQube.")
    add_number(doc, "Levantar el contenedor de SonarQube y exponer el puerto 9000.")
    add_number(doc, "Esperar a que el servicio quede disponible localmente.")
    add_command_block(
        doc,
        [
            "cd /Users/greivin/Documents/GitHub/Edu360/edu360_web",
            "docker compose up -d",
        ],
    )

    add_paragraph(
        doc,
        "El archivo de configuración utiliza la imagen sonarqube:community y publica el servicio en http://127.0.0.1:9000.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )

    add_heading(doc, "3. Acceso inicial a SonarQube", level=1)
    add_number(doc, "Abrir el navegador e ingresar a la dirección http://127.0.0.1:9000.")
    add_number(doc, "Esperar a que cargue la pantalla principal de SonarQube.")
    add_number(doc, "Iniciar sesión con una cuenta administradora local.")
    add_number(doc, "Una vez dentro, confirmar que la interfaz muestre el menú principal de proyectos.")

    add_heading(doc, "4. Creación del proyecto y preparación del análisis", level=1)
    add_paragraph(
        doc,
        "Después de acceder a SonarQube, se creó el proyecto de análisis y se generó el token correspondiente. El proyecto utilizado en esta práctica fue:",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_bullet(doc, "Clave del proyecto: mi_nextjs_app")
    add_bullet(doc, "Nombre del proyecto: Mi NextJS App")

    add_paragraph(
        doc,
        "Además, en la carpeta edu360_web se preparó el archivo sonar-project.properties con la configuración de nombre, clave del proyecto, rutas de código fuente y exclusiones necesarias.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )

    add_heading(doc, "5. Ejecución del análisis", level=1)
    add_paragraph(
        doc,
        "Con el proyecto ya configurado, se ejecutó SonarScanner desde la carpeta del frontend. El proyecto ya contaba con un script npm llamado sonar, por lo que el análisis se lanzó con el siguiente comando:",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_command_block(
        doc,
        [
            "cd /Users/greivin/Documents/GitHub/Edu360/edu360_web",
            "SONAR_HOST_URL=http://127.0.0.1:9000 SONAR_TOKEN=<token> npm run sonar",
        ],
    )
    add_paragraph(
        doc,
        "El proceso terminó con el mensaje ANALYSIS SUCCESSFUL, confirmando que el reporte fue enviado correctamente al servidor local de SonarQube.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )

    add_heading(doc, "6. Pasos para acceder al dashboard del proyecto", level=1)
    add_number(doc, "Ingresar a SonarQube en http://127.0.0.1:9000.")
    add_number(doc, "Ir a la sección Projects del menú principal.")
    add_number(doc, "Seleccionar el proyecto Mi NextJS App.")
    add_number(doc, "Abrir la vista Overview para ver el resumen general del análisis.")
    add_number(doc, "Si se desea ver detalles, abrir las secciones Issues, Security Hotspots, Measures o Activity.")

    add_paragraph(
        doc,
        "También es posible acceder directamente al dashboard del proyecto mediante la siguiente dirección:",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_command_block(
        doc,
        [
            "http://127.0.0.1:9000/dashboard?id=mi_nextjs_app",
        ],
    )

    add_heading(doc, "7. Observaciones finales", level=1)
    add_bullet(doc, "La instancia analizada está configurada en MQR Mode.")
    add_bullet(doc, "En ese modo, los hallazgos se visualizan como Reliability, Security y Maintainability.")
    add_bullet(doc, "El dashboard del proyecto queda disponible una vez que SonarQube procesa el análisis enviado por SonarScanner.")

    add_paragraph(
        doc,
        "Con estos pasos se completó la instalación local de SonarQube, la puesta en marcha del servicio y el acceso correcto al dashboard del proyecto para revisar los resultados del análisis estático.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_before=6,
    )

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Edu360 | Guía de instalación y acceso a SonarQube")
    style_run(r, size=9, color=MUTED)

    doc.save(path)


if __name__ == "__main__":
    build_doc("/Users/greivin/Documents/GitHub/Edu360/Entregable_SonarQube_Edu360.docx")
